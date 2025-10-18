from flask import Flask, request, jsonify
from werkzeug.utils import secure_filename
import csv
import openpyxl
import sqlite3
import os

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'csv', 'xlsx', 'xls'}
DATABASE = 'products.db'


if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)


def setupDb():
    conn = sqlite3.connect(DATABASE)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sku TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            brand TEXT NOT NULL,
            color TEXT,
            size TEXT,
            mrp REAL NOT NULL,
            price REAL NOT NULL,
            quantity INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    conn.commit()
    conn.close()


setupDb()


def isFileAllowed(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def checkProduct(row):
    errors = []
    required = ['sku', 'name', 'brand', 'mrp', 'price']

    for field in required:
        if field not in row or not str(row.get(field, '')).strip():
            errors.append(f'Missing required field: {field}')

    if errors:
        return {'isValid': False, 'errors': errors}

    try:
        mrp = float(row.get('mrp', 0))
        price = float(row.get('price', 0))
        quantity = int(row.get('quantity', 0))
    except (ValueError, TypeError):
        errors.append('Invalid numeric values for mrp, price, or quantity')
        return {'isValid': False, 'errors': errors}

    if mrp < 0:
        errors.append('Invalid MRP: must be a positive number')

    if price < 0:
        errors.append('Invalid price: must be a positive number')

    if price > mrp:
        errors.append(f'Price ({price}) must be ≤ MRP ({mrp})')

    if quantity < 0:
        errors.append('Invalid quantity: must be a non-negative number')

    if not str(row.get('sku', '')).strip():
        errors.append('SKU cannot be empty')

    return {
        'isValid': len(errors) == 0,
        'errors': errors
    }


def getCsvData(file_path):
    records = []
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                trimmed_row = {k: v.strip() if isinstance(v, str) else v for k, v in row.items()}
                records.append(trimmed_row)
    except Exception as e:
        raise Exception(f'CSV parsing error: {str(e)}')

    return records


def getExcelData(file_path):
    records = []
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active

        headers = []
        for cell in ws[1]:
            headers.append(cell.value)

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for i, header in enumerate(headers):
                value = row[i] if i < len(row) else None
                if isinstance(value, str):
                    row_dict[header] = value.strip()
                elif value is not None:
                    row_dict[header] = str(value).strip()
                else:
                    row_dict[header] = ''
            records.append(row_dict)
    except Exception as e:
        raise Exception(f'Excel parsing error: {str(e)}')

    return records


@app.route('/upload', methods=['POST'])
def uploadMyFile():
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file uploaded'}), 400

        file = request.files['file']

        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400

        if not isFileAllowed(file.filename):
            return jsonify({'error': 'Only CSV and Excel files are allowed'}), 400

        filename = secure_filename(file.filename)
        file_path = os.path.join(UPLOAD_FOLDER, filename)
        file.save(file_path)

        file_ext = filename.rsplit('.', 1)[1].lower()

        if file_ext == 'csv':
            records = getCsvData(file_path)
            file_type = 'CSV'
        else:
            records = getExcelData(file_path)
            file_type = 'Excel'

        failed = []
        stored = 0
        conn = sqlite3.connect(DATABASE)
        c = conn.cursor()

        for idx, row in enumerate(records, start=2):
            validation = checkProduct(row)

            if validation['isValid']:
                try:
                    c.execute('''
                        INSERT INTO products (sku, name, brand, color, size, mrp, price, quantity)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        row.get('sku', '').strip(),
                        row.get('name', '').strip(),
                        row.get('brand', '').strip(),
                        row.get('color', '').strip() or None,
                        row.get('size', '').strip() or None,
                        float(row.get('mrp', 0)),
                        float(row.get('price', 0)),
                        int(row.get('quantity', 0))
                    ))
                    stored += 1
                except sqlite3.IntegrityError as e:
                    failed.append({
                        'row': idx,
                        'sku': row.get('sku', 'unknown'),
                        'error': str(e)
                    })
            else:
                failed.append({
                    'row': idx,
                    'sku': row.get('sku', 'unknown'),
                    'errors': validation['errors']
                })

        conn.commit()
        conn.close()

        os.remove(file_path)

        return jsonify({
            'stored': stored,
            'failed': failed,
            'total': len(records),
            'fileType': file_type
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/products', methods=['GET'])
def showAllProducts():
    try:
        page = request.args.get('page', 1, type=int)
        limit = request.args.get('limit', 10, type=int)
        offset = (page - 1) * limit

        conn = sqlite3.connect(DATABASE)
        conn.row_factory = sqlite3.Row
        c = conn.cursor()

        c.execute('''
            SELECT * FROM products ORDER BY created_at DESC LIMIT ? OFFSET ?
        ''', (limit, offset))
        rows = c.fetchall()

        c.execute('SELECT COUNT(*) as total FROM products')
        total = c.fetchone()['total']

        conn.close()

        products = [dict(row) for row in rows]

        return jsonify({
            'data': products,
            'pagination': {
                'page': page,
                'limit': limit,
                'total': total,
                'pages': (total + limit - 1) // limit
            }
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/products/search', methods=['GET'])
def searchForProducts():
    try:
        brand = request.args.get('brand')
        color = request.args.get('color')
        min_price = request.args.get('minPrice', type=float)
        max_price = request.args.get('maxPrice', type=float)
        page = request.args.get('page', 1, type=int)
        limit = request.args.get('limit', 10, type=int)
        offset = (page - 1) * limit

        conn = sqlite3.connect(DATABASE)
        conn.row_factory = sqlite3.Row
        c = conn.cursor()

        query = 'SELECT * FROM products WHERE 1=1'
        params = []

        if brand:
            query += ' AND brand = ?'
            params.append(brand)

        if color:
            query += ' AND color = ?'
            params.append(color)

        if min_price is not None:
            query += ' AND price >= ?'
            params.append(min_price)

        if max_price is not None:
            query += ' AND price <= ?'
            params.append(max_price)

        query += ' ORDER BY created_at DESC LIMIT ? OFFSET ?'
        params.extend([limit, offset])

        c.execute(query, params)
        rows = c.fetchall()

        count_query = 'SELECT COUNT(*) as total FROM products WHERE 1=1'
        count_params = []

        if brand:
            count_query += ' AND brand = ?'
            count_params.append(brand)

        if color:
            count_query += ' AND color = ?'
            count_params.append(color)

        if min_price is not None:
            count_query += ' AND price >= ?'
            count_params.append(min_price)

        if max_price is not None:
            count_query += ' AND price <= ?'
            count_params.append(max_price)

        c.execute(count_query, count_params)
        total = c.fetchone()['total']

        conn.close()

        products = [dict(row) for row in rows]

        return jsonify({
            'data': products,
            'filters': {
                'brand': brand,
                'color': color,
                'minPrice': min_price,
                'maxPrice': max_price
            },
            'pagination': {
                'page': page,
                'limit': limit,
                'total': total,
                'pages': (total + limit - 1) // limit
            }
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/health', methods=['GET'])
def healthCheck():
    return jsonify({'status': 'ok'}), 200


@app.errorhandler(404)
def handle404(error):
    return jsonify({'error': 'Endpoint not found'}), 404


@app.errorhandler(500)
def handle500(error):
    return jsonify({'error': 'Internal server error'}), 500


if __name__ == '__main__':
    print('Product Catalog API running on http://localhost:5000')
    print('Endpoints:')
    print('  POST /upload - Upload CSV/Excel file')
    print('  GET /products - List all products')
    print('  GET /products/search - Search/filter products')
    print('  GET /health - Health check')
    app.run(debug=True, port=5000)